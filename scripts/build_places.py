# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl"]
# ///
"""Build the globe's Lakshadweep place list from the dataset workbook.

Reads the "Islands" sheet of data/source/BlueGrid_Lakshadweep_Dummy_Data.xlsx and
writes frontend/assets/data/places/lakshadweep.json.

Islands sharing an atoll/reef (and lying close together) are grouped under one
atoll place, so the Lakshadweep view shows one ring per atoll; clicking it zooms
in to show that atoll's islands. An atoll with a single island is just the island.

Run from the repo root:  uv run scripts/build_places.py
"""

import json
import math
import re
from collections import OrderedDict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "data/source/BlueGrid_Lakshadweep_Dummy_Data.xlsx"
OUTPUT = ROOT / "frontend/assets/data/places/lakshadweep.json"

REGION_ID = "IND-LD"  # the Lakshadweep click area defined in frontend/js/config.js
MAX_GROUP_SPREAD_KM = 20  # islands further apart than this aren't grouped
MIN_GROUP_ZOOM_M = 15_000  # closest camera height for an atoll view


def slug(text):
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")


def distance_km(a, b):
    """Great-circle distance between two (lat, lng) points."""
    lat1, lng1, lat2, lng2 = map(math.radians, (*a, *b))
    h = math.sin((lat2 - lat1) / 2) ** 2 + math.cos(lat1) * math.cos(lat2) * math.sin((lng2 - lng1) / 2) ** 2
    return 2 * 6371 * math.asin(math.sqrt(h))


def group_name(atoll):
    return atoll if re.search(r"\b(atoll|reef|bank|par)\b", atoll, re.I) else f"{atoll} Atoll"


def read_islands():
    sheet = load_workbook(SOURCE, read_only=True, data_only=True)["Islands"]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    for row in rows:
        if not any(row):  # the table ends at the first blank row; notes follow below it
            break
        record = dict(zip(header, row))
        yield {
            "name": record["island"],
            "atoll": record["atoll_or_reef"],
            "lat": round(float(record["lat_deg_N"]), 4),
            "lng": round(float(record["lon_deg_E"]), 4),
        }


def build_places(islands):
    atolls = OrderedDict()
    for island in islands:
        atolls.setdefault(island["atoll"], []).append(island)

    places = []
    for atoll, members in atolls.items():
        points = [(m["lat"], m["lng"]) for m in members]
        spread = max((distance_km(a, b) for a in points for b in points), default=0)

        if len(members) == 1 or spread > MAX_GROUP_SPREAD_KM:
            parent = REGION_ID
        else:
            name = group_name(atoll)
            parent = slug(name)
            places.append({
                "id": parent,
                "name": name,
                "lat": round(sum(p[0] for p in points) / len(points), 4),
                "lng": round(sum(p[1] for p in points) / len(points), 4),
                "parent": REGION_ID,
                "zoomHeight": max(MIN_GROUP_ZOOM_M, int(round(spread * 2500, -3))),
            })

        places += [
            {"id": slug(m["name"]), "name": m["name"], "lat": m["lat"], "lng": m["lng"], "parent": parent}
            for m in members
        ]

    ids = [p["id"] for p in places]
    duplicates = {i for i in ids if ids.count(i) > 1}
    if duplicates:
        raise SystemExit(f"Duplicate place ids: {sorted(duplicates)}")
    return places


def main():
    places = build_places(read_islands())
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps({
        "_generated": f"by scripts/build_places.py from {SOURCE.relative_to(ROOT)}; do not edit by hand",
        "places": places,
    }, indent=2, ensure_ascii=False) + "\n")
    groups = sum(1 for p in places if p["parent"] == REGION_ID)
    print(f"Wrote {len(places)} places ({groups} shown on the Lakshadweep view) to {OUTPUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
